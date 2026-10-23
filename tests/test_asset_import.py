# -*- coding: utf-8 -*-
"""资产台账批量导入（`asset_import` + CLI `assets-import` + 预览对话框文案）。

**它解决什么**：`assets` 那 6 个人工台账字段（采购日期 / 保修到期 / 负责人 /
序列号 / 位置树 / 自定义字段）此前只能一台一台填 —— 界面点「编辑资产信息…」，
或者命令行 `assets --set`。设备上百台时，这个字段最终就是空的。

覆盖点（都围绕"静默改错数据"这类失效）：

1. **空单元格 = 跳过，不是清空** —— 本模块最要紧的一条。把空当清空，
   一次导入就能静默抹掉一大片台账，且无法撤销。
2. **认设备只按身份键 / MAC** —— IP 会被 DHCP 换给别的设备，按 IP 认可能把
   A 的负责人写到 B 头上；要按 IP 必须显式 `--allow-ip`。
3. **先编译计划再写** —— 一行的坏日期不能变成"写了一半"。
4. **表头别名表不翻译** —— 它拿比较，不拿显示。
5. **写的是生命周期列，扫描列一个都不动**。

⚠ `cli.main()` 会 `i18n.init(None)` 切走**整个进程**的语言（CI 的 locale 是
en-US），所以每次调用都必须保存并还原 —— 照 `tests/test_cli_arp_mac.py`。
"""
import contextlib
import io
import json
import os
import shutil
import sys
import tempfile
import unittest
from unittest import mock

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import asset_import                                     # noqa: E402
import cli                                              # noqa: E402
import db_util                                          # noqa: E402
import history_store                                    # noqa: E402
import i18n                                             # noqa: E402
from ui_dialogs import asset_import_preview              # noqa: E402

#: 三台设备：一台普通 PC、一台随机 MAC 手机、一台交换机。
DEVS = [
    {"ip": "10.0.0.10", "status": "online", "mac": "A8:4A:28:63:1A:46",
     "hostname": "pc-a", "vendor": "Hon Hai", "device_class": "台式机"},
    {"ip": "10.0.0.11", "status": "online", "mac": "E6:56:23:5E:A3:0C",
     "hostname": "phone-b", "vendor": "", "device_class": "手机"},
    {"ip": "10.0.0.12", "status": "online", "mac": "64:C3:94:88:FB:C5",
     "hostname": "sw-c", "vendor": "Tenda", "device_class": "网络设备"},
]

KEY_PC = "a8:4a:28:63:1a:46"
KEY_PHONE = "e6:56:23:5e:a3:0c"
KEY_SW = "64:c3:94:88:fb:c5"


class _DbCase(unittest.TestCase):
    """每个用例一个独立临时库（三台设备）。"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="ipsc_import_")
        self.db = os.path.join(self.tmp, "h.db")
        history_store.save_snapshot(DEVS, note="fixture", db_path=self.db)
        self.assets = history_store.list_assets(db_path=self.db)

    def tearDown(self):
        db_util.close_connection(self.db)
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _plan(self, rows, **kwargs):
        return asset_import.build_plan(rows, self.assets, **kwargs)

    def _reload(self):
        """**关连接后**重新读 —— 验"真的落盘了"，不是只改了缓存。"""
        db_util.close_connection(self.db)
        return {r["identity_key"]: r
                for r in history_store.list_assets(db_path=self.db)}

    def _write_csv(self, text, name="t.csv", encoding="utf-8"):
        path = os.path.join(self.tmp, name)
        with open(path, "w", encoding=encoding, newline="") as f:
            f.write(text)
        return path


# ══════════════════════════════════════════════════════════════════════
#  MAC 归一化的重复实现：必须与另外两处逐字节一致
# ══════════════════════════════════════════════════════════════════════

class MacKeyConsistencyTest(unittest.TestCase):
    """`asset_import._mac_key` 是**刻意复制**的第三份实现（见其 docstring）。

    三处口径一旦漂移，就会出现「地址库认得、导入认不得」这种最难查的问题。
    """

    SAMPLES = [
        "A8:4A:28:63:1A:46", "a8-4a-28-63-1a-46", "a84a28631a46",
        "A8 4A 28 63 1A 46", "  e6:56:23:5e:a3:0c  ",
    ]

    def test_agrees_with_excel_handler(self):
        import excel_handler
        for s in self.SAMPLES:
            mine = asset_import._mac_key(s)
            theirs = excel_handler.normalize_mac(s).replace(":", "")
            self.assertEqual(mine, theirs, msg="样本 %r 上两份实现不一致" % s)

    def test_agrees_with_topology_overrides(self):
        import topology_overrides
        for s in self.SAMPLES:
            self.assertEqual(asset_import._mac_key(s),
                             topology_overrides.normalize_mac(s),
                             msg="样本 %r 上两份实现不一致" % s)

    def test_rejects_non_mac(self):
        for s in ("", None, "zz:zz:zz:zz:zz:zz", "1", "0x02"):
            self.assertEqual(asset_import._mac_key(s), "", msg=repr(s))


# ══════════════════════════════════════════════════════════════════════
#  单元格 / 表头
# ══════════════════════════════════════════════════════════════════════

class CellTextTest(unittest.TestCase):
    def test_excel_float_int_becomes_int(self):
        # Excel 把整数读成 float：序列号 12345 不能变成 "12345.0"。
        self.assertEqual(asset_import._cell_text(12345.0), "12345")
        self.assertEqual(asset_import._cell_text(12345), "12345")
        self.assertEqual(asset_import._cell_text(1.5), "1.5")

    def test_datetime_becomes_iso_date(self):
        import datetime
        self.assertEqual(
            asset_import._cell_text(datetime.datetime(2026, 5, 1, 0, 0)),
            "2026-05-01")
        self.assertEqual(asset_import._cell_text(datetime.date(2026, 5, 1)),
                         "2026-05-01")

    def test_bool_is_not_int(self):
        # bool 是 int 的子类：判断顺序写反会得到 "1" / "0"。
        self.assertEqual(asset_import._cell_text(True), "true")
        self.assertEqual(asset_import._cell_text(False), "false")

    def test_none_and_whitespace(self):
        self.assertEqual(asset_import._cell_text(None), "")
        self.assertEqual(asset_import._cell_text("  x  "), "x")


class DetectColumnsTest(unittest.TestCase):
    def test_english_headers(self):
        mapping, ignored = asset_import.detect_columns(
            ["identity_key", "mac", "purchase_date", "owner"])
        self.assertEqual(mapping, {"identity_key": 0, "mac": 1,
                                   "purchase_date": 2, "owner": 3})
        self.assertEqual(ignored, [])

    def test_chinese_headers(self):
        mapping, _ = asset_import.detect_columns(
            ["身份键", "采购日期", "保修到期", "负责人", "序列号",
             "位置树", "自定义字段"])
        self.assertEqual(mapping["identity_key"], 0)
        self.assertEqual(mapping["purchase_date"], 1)
        self.assertEqual(mapping["warranty_end"], 2)
        self.assertEqual(mapping["owner"], 3)
        self.assertEqual(mapping["serial"], 4)
        self.assertEqual(mapping["location_tree"], 5)
        self.assertEqual(mapping["custom_json"], 6)

    def test_tolerates_case_space_star_and_bom(self):
        mapping, _ = asset_import.detect_columns(
            ["\ufeffOwner", " Purchase_Date *", "MAC 地址 "])
        self.assertEqual(mapping, {"owner": 0, "purchase_date": 1, "mac": 2})

    def test_unknown_headers_are_reported_not_dropped_silently(self):
        _, ignored = asset_import.detect_columns(["owner", "备注", "位置"])
        self.assertEqual(ignored, ["备注", "位置"])

    def test_dangerous_aliases_are_not_mapped(self):
        """`位置` / `使用者` 是**扫描写出来的列**，不能当人工台账写进去。

        `location` / `user` 由扫描刷新，`location_tree` / `owner` 是人填的 ——
        把前者当后者导入，等于用扫描值覆盖用户手填的台账。
        """
        mapping, ignored = asset_import.detect_columns(["位置", "使用者", "location", "user"])
        self.assertEqual(mapping, {})
        self.assertEqual(ignored, ["位置", "使用者", "location", "user"])

    def test_duplicate_column_keeps_first(self):
        mapping, _ = asset_import.detect_columns(["owner", "负责人"])
        self.assertEqual(mapping, {"owner": 0})


# ══════════════════════════════════════════════════════════════════════
#  计划编译
# ══════════════════════════════════════════════════════════════════════

class BuildPlanTest(_DbCase):
    def test_ok_row(self):
        plan = self._plan([["identity_key", "owner"], [KEY_PC, "张三"]])
        self.assertEqual(plan["problem"], "")
        item = plan["items"][0]
        self.assertEqual(item["status"], asset_import.ST_OK)
        self.assertEqual(item["identity_key"], KEY_PC)
        self.assertEqual(item["fields"], {"owner": "张三"})
        self.assertEqual(item["line"], 2)

    def test_empty_cell_is_skipped_not_cleared(self):
        """🔴 空单元格默认**跳过** —— 这条错了会静默抹掉已有台账。"""
        history_store.update_asset_lifecycle(KEY_PC, db_path=self.db, owner="张三")
        self.assets = history_store.list_assets(db_path=self.db)
        plan = self._plan([["identity_key", "owner", "serial"], [KEY_PC, "", "SN-9"]])
        self.assertEqual(plan["items"][0]["fields"], {"serial": "SN-9"})
        asset_import.apply_plan(plan, db_path=self.db)
        self.assertEqual(self._reload()[KEY_PC]["owner"], "张三")

    def test_clear_empty_clears_the_field(self):
        history_store.update_asset_lifecycle(KEY_PC, db_path=self.db, owner="张三")
        self.assets = history_store.list_assets(db_path=self.db)
        plan = self._plan([["identity_key", "owner"], [KEY_PC, ""]], clear_empty=True)
        self.assertEqual(plan["items"][0]["status"], asset_import.ST_OK)
        asset_import.apply_plan(plan, db_path=self.db)
        self.assertEqual(self._reload()[KEY_PC]["owner"], "")

    def test_date_is_normalized(self):
        plan = self._plan([["identity_key", "purchase_date"], [KEY_PC, "2026/5/1"]])
        self.assertEqual(plan["items"][0]["fields"]["purchase_date"], "2026-05-01")

    def test_bad_date_is_invalid_and_names_the_field(self):
        plan = self._plan([["identity_key", "purchase_date"], [KEY_PC, "2026-13-45"]])
        item = plan["items"][0]
        self.assertEqual(item["status"], asset_import.ST_INVALID)
        self.assertEqual(item["reason"], asset_import.REASON_BAD_DATE)
        self.assertEqual(item["field"], "purchase_date")
        self.assertEqual(item["detail"], "2026-13-45")

    def test_bad_custom_json_is_invalid(self):
        plan = self._plan([["identity_key", "custom_json"], [KEY_PC, "{不是JSON"]])
        item = plan["items"][0]
        self.assertEqual(item["status"], asset_import.ST_INVALID)
        self.assertEqual(item["reason"], asset_import.REASON_BAD_JSON)

    def test_custom_json_keys_are_not_translated(self):
        plan = self._plan([["identity_key", "custom_json"],
                           [KEY_PC, '{"机架": "A3"}']])
        self.assertEqual(plan["items"][0]["fields"]["custom_json"],
                         '{"机架": "A3"}')

    def test_unknown_identity_key(self):
        plan = self._plan([["identity_key", "owner"], ["nope", "张三"]])
        item = plan["items"][0]
        self.assertEqual(item["status"], asset_import.ST_NOTFOUND)
        self.assertEqual(item["reason"], asset_import.REASON_UNKNOWN_DEVICE)

    def test_mac_match_ignores_separator_and_case(self):
        for mac in ("a8-4a-28-63-1a-46", "A84A28631A46", "a8:4a:28:63:1a:46"):
            plan = self._plan([["mac", "owner"], [mac, "张三"]])
            self.assertEqual(plan["items"][0]["identity_key"], KEY_PC, msg=mac)

    def test_ip_alone_is_refused_by_default(self):
        plan = self._plan([["ip", "owner"], ["10.0.0.10", "张三"]])
        item = plan["items"][0]
        self.assertEqual(item["status"], asset_import.ST_INVALID)
        self.assertEqual(item["reason"], asset_import.REASON_IP_NOT_ALLOWED)

    def test_ip_alone_works_with_allow_ip(self):
        plan = self._plan([["ip", "owner"], ["10.0.0.10", "张三"]], allow_ip=True)
        item = plan["items"][0]
        self.assertEqual(item["status"], asset_import.ST_OK)
        self.assertEqual(item["identity_key"], KEY_PC)

    def test_duplicate_rows_keep_first(self):
        plan = self._plan([["identity_key", "owner"],
                           [KEY_PC, "张三"], [KEY_PC, "李四"]])
        self.assertEqual(plan["items"][0]["status"], asset_import.ST_OK)
        self.assertEqual(plan["items"][1]["status"], asset_import.ST_DUPLICATE)
        self.assertEqual(plan["items"][1]["reason"], asset_import.REASON_DUPLICATE_ROW)
        self.assertEqual(plan["items"][1]["detail"], "2")   # 第一行在第 2 行

    def test_empty_row_is_not_an_error(self):
        plan = self._plan([["identity_key", "owner"], ["", ""]])
        self.assertEqual(plan["items"][0]["status"], asset_import.ST_EMPTY)

    def test_no_identity_at_all(self):
        """有身份键列、但这一格是空的，也没有 MAC/IP ⇒ 行级 `no_identity`。

        ⚠️ 整张表**没有**身份列是另一回事，走 `PROBLEM_NO_IDENTITY_COLUMN`
           （见 `PlanProblemTest`），不会编译出逐行计划。
        """
        plan = self._plan([["identity_key", "owner"], ["", "张三"]])
        item = plan["items"][0]
        self.assertEqual(item["status"], asset_import.ST_INVALID)
        self.assertEqual(item["reason"], asset_import.REASON_NO_IDENTITY)

    def test_identical_value_is_nochange(self):
        history_store.update_asset_lifecycle(KEY_PC, db_path=self.db, owner="张三")
        self.assets = history_store.list_assets(db_path=self.db)
        plan = self._plan([["identity_key", "owner"], [KEY_PC, "张三"]])
        item = plan["items"][0]
        self.assertEqual(item["status"], asset_import.ST_NOCHANGE)
        self.assertEqual(item["reason"], asset_import.REASON_ALREADY_EQUAL)

    def test_no_writable_column_present(self):
        plan = self._plan([["identity_key", "owner"], [KEY_PC, ""]])
        item = plan["items"][0]
        self.assertEqual(item["status"], asset_import.ST_NOCHANGE)
        self.assertEqual(item["reason"], asset_import.REASON_NO_FIELDS)


class PlanProblemTest(_DbCase):
    def test_empty_file(self):
        plan = asset_import.build_plan([], self.assets)
        self.assertEqual(plan["problem"], asset_import.PROBLEM_EMPTY)

    def test_no_writable_columns(self):
        plan = self._plan([["foo", "bar"], ["1", "2"]])
        self.assertEqual(plan["problem"], asset_import.PROBLEM_NO_COLUMNS)

    def test_no_identity_column(self):
        plan = self._plan([["owner"], ["张三"]])
        self.assertEqual(plan["problem"], asset_import.PROBLEM_NO_IDENTITY_COLUMN)

    def test_plan_does_not_mutate_inputs(self):
        rows = [["identity_key", "owner"], [KEY_PC, "张三"]]
        assets_before = [dict(r) for r in self.assets]
        self._plan(rows)
        self.assertEqual(rows, [["identity_key", "owner"], [KEY_PC, "张三"]])
        self.assertEqual(self.assets, assets_before)


# ══════════════════════════════════════════════════════════════════════
#  读文件
# ══════════════════════════════════════════════════════════════════════

class ReadTableTest(_DbCase):
    def test_csv_utf8_with_bom(self):
        path = self._write_csv("\ufeffidentity_key,owner\n%s,张三\n" % KEY_PC,
                               encoding="utf-8")
        headers, body = asset_import.read_table(path)
        # `utf-8-sig` 会把 BOM 吃掉；即便没吃掉，`detect_columns` 也会去 BOM。
        self.assertEqual(list(headers), ["identity_key", "owner"])
        self.assertEqual(len(body), 1)
        mapping, _ = asset_import.detect_columns(headers)
        self.assertIn("identity_key", mapping)

    def test_csv_gbk(self):
        path = self._write_csv("身份键,负责人\n%s,张三\n" % KEY_PC, encoding="gbk")
        headers, body = asset_import.read_table(path)
        self.assertEqual(list(headers), ["身份键", "负责人"])
        self.assertEqual(body[0][1], "张三")

    def test_tsv(self):
        path = self._write_csv("identity_key\towner\n%s\t张三\n" % KEY_PC,
                               name="t.tsv")
        headers, body = asset_import.read_table(path)
        self.assertEqual(list(headers), ["identity_key", "owner"])
        self.assertEqual(body[0][1], "张三")

    def test_xlsx(self):
        import openpyxl
        path = os.path.join(self.tmp, "t.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["identity_key", "owner", "purchase_date"])
        ws.append([KEY_PC, "张三", "2026-05-01"])
        wb.save(path)
        headers, body = asset_import.read_table(path)
        self.assertEqual(list(headers), ["identity_key", "owner", "purchase_date"])
        self.assertEqual(body[0][1], "张三")

    def test_missing_file_raises_value_error(self):
        with self.assertRaises(ValueError):
            asset_import.read_table(os.path.join(self.tmp, "nope.csv"))


# ══════════════════════════════════════════════════════════════════════
#  写入
# ══════════════════════════════════════════════════════════════════════

class ApplyPlanTest(_DbCase):
    def test_writes_persist_after_closing_the_connection(self):
        """🔴 必须关连接后新连接再读 —— 只读缓存连接会漏掉"没 commit"。"""
        plan = self._plan([["identity_key", "owner", "purchase_date"],
                           [KEY_PC, "张三", "2026-05-01"]])
        result = asset_import.apply_plan(plan, db_path=self.db)
        self.assertEqual(result["written"], 1)
        row = self._reload()[KEY_PC]
        self.assertEqual(row["owner"], "张三")
        self.assertEqual(row["purchase_date"], "2026-05-01")

    def test_scan_columns_are_untouched(self):
        """导入只碰生命周期列 —— 扫描写出来的列一个都不能变。"""
        before = self._reload()[KEY_PC]
        plan = self._plan([["identity_key", "owner"], [KEY_PC, "张三"]])
        asset_import.apply_plan(plan, db_path=self.db)
        after = self._reload()[KEY_PC]
        for col in ("ip", "mac", "hostname", "vendor", "device_class", "status"):
            self.assertEqual(after[col], before[col], msg=col)

    def test_apply_is_idempotent(self):
        plan = self._plan([["identity_key", "owner"], [KEY_PC, "张三"]])
        asset_import.apply_plan(plan, db_path=self.db)
        # 同一份计划再 apply 一次：状态已是 written，不会再写。
        again = asset_import.apply_plan(plan, db_path=self.db)
        self.assertEqual(again["written"], 0)
        # 重新读库后再编译，则会判成"无需改动"（`self.assets` 是 setUp 的快照，
        # 不重读的话比对的是写入之前的值 —— 那是测试自己的问题，不是代码的）。
        self.assets = history_store.list_assets(db_path=self.db)
        fresh = self._plan([["identity_key", "owner"], [KEY_PC, "张三"]])
        self.assertEqual(fresh["items"][0]["status"], asset_import.ST_NOCHANGE)

    def test_apply_skips_non_ok_rows(self):
        plan = self._plan([["identity_key", "owner"],
                           [KEY_PC, "张三"], ["nope", "李四"]])
        result = asset_import.apply_plan(plan, db_path=self.db)
        self.assertEqual(result["written"], 1)
        self.assertEqual(plan["items"][1]["status"], asset_import.ST_NOTFOUND)

    def test_missing_db_reports_notfound_not_success(self):
        """库层返回 0 行 ⇒ 判"找不到"，**不能**算写成功。"""
        plan = self._plan([["identity_key", "owner"], [KEY_PC, "张三"]])
        db_util.close_connection(self.db)
        os.remove(self.db)
        result = asset_import.apply_plan(plan, db_path=self.db)
        self.assertEqual(result["written"], 0)
        self.assertEqual(len(result["failed"]), 1)


# ══════════════════════════════════════════════════════════════════════
#  报告文案
# ══════════════════════════════════════════════════════════════════════

def _has_cjk(text) -> bool:
    return any("\u4e00" <= ch <= "\u9fff" for ch in text)


class ReportTest(_DbCase):
    def _report(self, plan, lang, written=None):
        old = i18n.current_lang()
        try:
            i18n.init(lang)
            return asset_import_preview.build_report(plan, written=written)
        finally:
            i18n.init(old)

    def test_report_lists_rows_needing_attention(self):
        plan = self._plan([["identity_key", "owner", "purchase_date"],
                           [KEY_PC, "张三", "2026-13-45"]])
        text = self._report(plan, "zh-CN")
        self.assertIn(KEY_PC, text)
        self.assertIn("2026-13-45", text)
        self.assertIn("第 2 行", text)

    def test_english_report_has_no_chinese(self):
        """英文界面下报告里不能漏中文 —— 三张标签表必须过 `tr()`。"""
        plan = self._plan([["identity_key", "owner", "purchase_date"],
                           [KEY_PC, "张三", "2026-13-45"],
                           ["nope", "李四", ""]])
        text = self._report(plan, "en-US")
        self.assertFalse(_has_cjk(text), msg="英文报告里出现中文：\n" + text)

    def test_written_line_appears_only_after_apply(self):
        plan = self._plan([["identity_key", "owner"], [KEY_PC, "张三"]])
        preview = self._report(plan, "zh-CN")
        done = self._report(plan, "zh-CN", written=1)
        self.assertEqual(len(done.splitlines()), len(preview.splitlines()) + 1)

    def test_problem_text_is_translated(self):
        old = i18n.current_lang()
        try:
            i18n.init("en-US")
            self.assertFalse(_has_cjk(asset_import_preview.problem_text(
                asset_import.PROBLEM_NO_COLUMNS)))
            self.assertEqual(asset_import_preview.problem_text(""), "")
        finally:
            i18n.init(old)


class LabelPackTest(unittest.TestCase):
    """三张标签表里的中文串必须在语言包里 —— 与 `test_i18n_lazy` 的登记互补。"""

    def test_all_labels_have_translations(self):
        import io as _io
        with _io.open(os.path.join(ROOT, "i18n", "en-US.json"),
                      encoding="utf-8") as f:
            pack = json.load(f)
        missing = []
        for name, table in (
            ("_STATUS_TEXT", asset_import_preview._STATUS_TEXT),
            ("_REASON_TEXT", asset_import_preview._REASON_TEXT),
            ("_PROBLEM_TEXT", asset_import_preview._PROBLEM_TEXT),
        ):
            for value in table.values():
                if _has_cjk(value) and value not in pack:
                    missing.append("%s: %r" % (name, value))
        self.assertEqual(missing, [], msg="缺英文翻译：\n  " + "\n  ".join(missing))


# ══════════════════════════════════════════════════════════════════════
#  CLI
# ══════════════════════════════════════════════════════════════════════

class CliTest(_DbCase):
    def _run(self, argv):
        """运行 `cli.main`，返回 (exit_code, stdout, stderr)，并还原语言。"""
        out, err = io.StringIO(), io.StringIO()
        lang = i18n.current_lang()
        try:
            with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
                code = cli.main(argv)
        finally:
            i18n.init(lang)
        return code, out.getvalue(), err.getvalue()

    def test_dry_run_writes_nothing(self):
        path = self._write_csv("identity_key,owner\n%s,张三\n" % KEY_PC)
        code, out, _ = self._run(["assets-import", path, "--db", self.db, "--dry-run"])
        self.assertEqual(code, cli.EXIT_OK)
        self.assertIn("将写入", out)
        self.assertEqual(self._reload()[KEY_PC]["owner"], "")

    def test_apply_writes(self):
        path = self._write_csv("identity_key,owner\n%s,张三\n" % KEY_PC)
        code, out, _ = self._run(["assets-import", path, "--db", self.db])
        self.assertEqual(code, cli.EXIT_OK)
        self.assertIn("已写入 1 台", out)
        self.assertEqual(self._reload()[KEY_PC]["owner"], "张三")

    def test_bad_row_exits_nonzero(self):
        path = self._write_csv("identity_key,purchase_date\n%s,2026-13-45\n" % KEY_PC)
        code, out, _ = self._run(["assets-import", path, "--db", self.db])
        self.assertEqual(code, cli.EXIT_ERROR)
        self.assertIn("日期格式无效", out)

    def test_json_keys_are_english(self):
        path = self._write_csv("identity_key,owner\n%s,张三\n" % KEY_PC)
        code, out, _ = self._run(
            ["assets-import", path, "--db", self.db, "--dry-run", "-f", "json"])
        self.assertEqual(code, cli.EXIT_OK)
        payload = json.loads(out)
        self.assertEqual(
            set(payload),
            {"file", "dry_run", "total", "columns", "ignored_columns",
             "counts", "written", "items"})
        item = payload["items"][0]
        self.assertEqual(
            set(item),
            {"line", "status", "reason", "detail", "field",
             "identity_key", "fields"})

    def test_no_columns_exits_nonzero(self):
        path = self._write_csv("foo,bar\n1,2\n")
        code, out, _ = self._run(["assets-import", path, "--db", self.db])
        self.assertEqual(code, cli.EXIT_ERROR)
        self.assertIn("没有识别到任何可写入的列", out)

    def test_missing_file_exits_nonzero(self):
        code, out, _ = self._run(
            ["assets-import", os.path.join(self.tmp, "nope.csv"), "--db", self.db])
        self.assertEqual(code, cli.EXIT_ERROR)
        self.assertIn("读不了文件", out)

    def test_clear_empty_flag(self):
        history_store.update_asset_lifecycle(KEY_PC, db_path=self.db, owner="张三")
        path = self._write_csv("identity_key,owner\n%s,\n" % KEY_PC)
        code, _, _ = self._run(
            ["assets-import", path, "--db", self.db, "--clear-empty"])
        self.assertEqual(code, cli.EXIT_OK)
        self.assertEqual(self._reload()[KEY_PC]["owner"], "")

    def test_cli_call_does_not_leak_process_language(self):
        """`cli.main()` 会 `i18n.init(None)` —— 必须还原，否则后续模块全红。"""
        before = i18n.current_lang()
        path = self._write_csv("identity_key,owner\n%s,张三\n" % KEY_PC)
        self._run(["assets-import", path, "--db", self.db, "--dry-run"])
        self.assertEqual(i18n.current_lang(), before)


# ══════════════════════════════════════════════════════════════════════
#  边界分支（覆盖度：这些路径正常流程走不到，但一旦走错就是静默出错）
# ══════════════════════════════════════════════════════════════════════

class CellTextEdgeTest(unittest.TestCase):
    def test_other_object_falls_back_to_str(self):
        # 不是 str / 数字 / 日期的对象：`str()` 兜底，绝不抛。
        self.assertEqual(asset_import._cell_text(["a"]), "['a']")


class DetectColumnsEdgeTest(unittest.TestCase):
    def test_blank_header_is_skipped(self):
        mapping, ignored = asset_import.detect_columns(["owner", "", None, "  "])
        self.assertEqual(mapping, {"owner": 0})
        self.assertEqual(ignored, [])


class ReadTableEdgeTest(_DbCase):
    def test_empty_csv(self):
        path = self._write_csv("")
        self.assertEqual(asset_import.read_table(path), ((), []))

    def test_single_column_csv_uses_fallback_delimiter(self):
        """只有一列时 `csv.Sniffer` 猜不出分隔符 —— 必须兜底，不能抛。"""
        path = self._write_csv("owner\n张三\n")
        headers, body = asset_import.read_table(path)
        self.assertEqual(list(headers), ["owner"])
        self.assertEqual(body[0][0], "张三")

    def test_undecodable_file_raises_value_error(self):
        path = os.path.join(self.tmp, "bad.csv")
        with open(path, "wb") as f:
            f.write(b"\xff\xfe\xfd\xfc")
        with self.assertRaises(ValueError):
            asset_import.read_table(path)

    def test_empty_xlsx(self):
        import openpyxl
        path = os.path.join(self.tmp, "empty.xlsx")
        openpyxl.Workbook().save(path)
        self.assertEqual(asset_import.read_table(path), ((), []))


class BuildPlanEdgeTest(_DbCase):
    def test_unknown_mac(self):
        plan = self._plan([["mac", "owner"], ["00:11:22:33:44:55", "张三"]])
        item = plan["items"][0]
        self.assertEqual(item["status"], asset_import.ST_NOTFOUND)
        self.assertEqual(item["reason"], asset_import.REASON_UNKNOWN_DEVICE)
        self.assertEqual(item["detail"], "00:11:22:33:44:55")

    def test_unknown_ip_with_allow_ip(self):
        plan = self._plan([["ip", "owner"], ["10.9.9.9", "张三"]], allow_ip=True)
        item = plan["items"][0]
        self.assertEqual(item["status"], asset_import.ST_NOTFOUND)

    def test_index_assets_skips_junk(self):
        """`assets` 里混进非字典 / 空身份键的行时不能炸，也不能把它们索引进去。"""
        by_key, by_mac, _ = asset_import._index_assets(
            [None, "x", {"identity_key": ""}, {"identity_key": "k", "mac": "zz"}])
        self.assertEqual(by_key, {"k": {"identity_key": "k", "mac": "zz"}})
        self.assertEqual(by_mac, {})


class PlanHelperTest(_DbCase):
    def test_pending_count_and_summarize(self):
        plan = self._plan([["identity_key", "owner"],
                           [KEY_PC, "张三"], ["nope", "李四"], ["", ""]])
        self.assertEqual(asset_import.pending_count(plan), 1)
        counts = asset_import.summarize(plan)
        self.assertEqual(counts[asset_import.ST_OK], 1)
        self.assertEqual(counts[asset_import.ST_NOTFOUND], 1)
        self.assertEqual(counts[asset_import.ST_EMPTY], 1)

    def test_pending_count_tolerates_empty_plan(self):
        self.assertEqual(asset_import.pending_count(None), 0)
        self.assertEqual(asset_import.summarize(None), {})


class ApplyPlanEdgeTest(_DbCase):
    def test_write_rejected_by_store_is_reported_per_row(self):
        """库层拒绝（`ValueError`）只影响那一行，其余行继续写。"""
        plan = self._plan([["identity_key", "owner"],
                           [KEY_PC, "张三"], [KEY_PHONE, "李四"]])
        real = history_store.update_asset_lifecycle

        def fake(key, db_path=None, **fields):
            if key == KEY_PC:
                raise ValueError("字段被拒绝")
            return real(key, db_path=db_path, **fields)

        with mock.patch("history_store.update_asset_lifecycle", side_effect=fake):
            result = asset_import.apply_plan(plan, db_path=self.db)
        self.assertEqual(result["written"], 1)
        self.assertEqual(len(result["failed"]), 1)
        self.assertEqual(result["failed"][0]["status"], asset_import.ST_INVALID)
        self.assertEqual(result["failed"][0]["reason"],
                         asset_import.REASON_WRITE_FAILED)
        self.assertEqual(result["failed"][0]["detail"], "字段被拒绝")
        self.assertEqual(self._reload()[KEY_PHONE]["owner"], "李四")


# ══════════════════════════════════════════════════════════════════════
#  模板
# ══════════════════════════════════════════════════════════════════════

class TemplateTest(_DbCase):
    """模板必须**只含导入真正会读的列**，且原样导回去不产生任何写入。

    后者是这张表的核心承诺：模板 → 填值 → 导入，未填的格子必须一个都不动。
    """

    def _run(self, argv):
        out, err = io.StringIO(), io.StringIO()
        lang = i18n.current_lang()
        try:
            with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
                code = cli.main(argv)
        finally:
            i18n.init(lang)
        return code, out.getvalue(), err.getvalue()

    def test_columns_are_exactly_the_nine(self):
        """模板只含 3 个认设备的列 + 6 个台账列 —— 扫描写入的 16 列一个都不许进来。"""
        self.assertEqual(asset_import.TEMPLATE_COLUMNS, (
            "identity_key", "mac", "ip", "purchase_date", "warranty_end",
            "owner", "serial", "location_tree", "custom_json"))

    def test_header_only_when_no_assets(self):
        rows = asset_import.build_template_rows()
        self.assertEqual(len(rows), 1)
        self.assertEqual(tuple(rows[0]), asset_import.TEMPLATE_COLUMNS)

    def test_fills_identity_but_leaves_lifecycle_blank(self):
        rows = asset_import.build_template_rows(self.assets)
        self.assertEqual(len(rows), 1 + len(self.assets))
        for row in rows[1:]:
            self.assertTrue(row[0], msg="identity_key 必须填好（用户编不出来）")
            self.assertEqual(row[3:], [""] * 6, msg="台账列必须留空")

    def test_junk_assets_are_skipped(self):
        rows = asset_import.build_template_rows([None, "x", {"identity_key": "k"}])
        self.assertEqual(len(rows), 2)

    def test_csv_roundtrip_writes_nothing(self):
        path = os.path.join(self.tmp, "tpl.csv")
        n = asset_import.write_template(path, assets=self.assets)
        self.assertEqual(n, 3)
        headers, body = asset_import.read_table(path)
        self.assertEqual(tuple(headers), asset_import.TEMPLATE_COLUMNS)
        plan = asset_import.build_plan([headers] + list(body), self.assets)
        self.assertEqual(plan["problem"], "")
        for item in plan["items"]:
            self.assertEqual(item["status"], asset_import.ST_NOCHANGE)

    def test_csv_ignores_notes(self):
        """CSV 装不下第二个 sheet —— notes 必须被丢弃，不能变成一行脏数据。"""
        path = os.path.join(self.tmp, "tpl2.csv")
        asset_import.write_template(path, assets=self.assets, notes=["说明行"])
        headers, body = asset_import.read_table(path)
        self.assertEqual(len(body), len(self.assets))

    def test_xlsx_has_a_notes_sheet(self):
        import openpyxl
        path = os.path.join(self.tmp, "tpl.xlsx")
        n = asset_import.write_template(
            path, assets=self.assets, notes=["第一行", "", "第三行"],
            sheet_title="台账", notes_title="说明")
        self.assertEqual(n, len(self.assets))
        wb = openpyxl.load_workbook(path)
        try:
            self.assertEqual(wb.sheetnames, ["台账", "说明"])
            ws = wb["台账"]
            self.assertEqual(ws.max_column, len(asset_import.TEMPLATE_COLUMNS))
            self.assertEqual(ws.freeze_panes, "A2")
            ns = wb["说明"]
            self.assertEqual(ns.cell(row=1, column=1).value, "第一行")
            self.assertEqual(ns.cell(row=3, column=1).value, "第三行")
        finally:
            wb.close()

    def test_unsupported_extension_raises(self):
        with self.assertRaises(ValueError):
            asset_import.write_template(os.path.join(self.tmp, "tpl.pdf"))

    def test_cli_writes_template(self):
        path = os.path.join(self.tmp, "cli_tpl.csv")
        code, out, _ = self._run(["assets-template", path, "--db", self.db])
        self.assertEqual(code, cli.EXIT_OK)
        self.assertIn("3 台设备待填", out)
        headers, body = asset_import.read_table(path)
        self.assertEqual(len(body), 3)

    def test_cli_empty_flag_writes_header_only(self):
        path = os.path.join(self.tmp, "cli_empty.csv")
        code, out, _ = self._run(
            ["assets-template", path, "--db", self.db, "--empty"])
        self.assertEqual(code, cli.EXIT_OK)
        self.assertIn("--empty", out)
        headers, body = asset_import.read_table(path)
        self.assertEqual(tuple(headers), asset_import.TEMPLATE_COLUMNS)
        self.assertEqual(body, [])

    def test_cli_bad_extension_exits_nonzero(self):
        code, out, _ = self._run(
            ["assets-template", os.path.join(self.tmp, "x.pdf"), "--db", self.db])
        self.assertEqual(code, cli.EXIT_ERROR)
        self.assertIn("只支持", out)


if __name__ == "__main__":
    unittest.main()
