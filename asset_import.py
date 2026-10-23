# -*- coding: utf-8 -*-
"""资产台账批量导入（CSV / Excel）的**纯逻辑**部分。

为什么单独开一个模块
--------------------
`asset_lifecycle.py` 管「字段怎么判定」，`history_store.py` 管「怎么存」，
本模块管「一张表怎么变成一串写入」。拆法与 `asset_lifecycle` 一致：

- **可测**：输入是「二维行数据 + 资产列表」，输出是「导入计划」——
  不建库、不需要 Qt，测试是纯函数级的。
- **两条入口共用同一套判定**：GUI 菜单与 CLI 子命令走**同一份**计划编译，
  否则「界面说 12 台、命令行说 13 台」这种矛盾迟早出现。

🔴 本模块**不 import `i18n`**、**不 import PyQt5**、**顶层不 import sqlite**
   （`apply_plan()` 里才惰性 import `history_store`）。
   产出的状态与原因一律是**语言无关的英文码**（`ok` / `notfound` / `bad_date`…），
   显示文字由调用方翻译 —— 与 `asset_lifecycle.warranty_state()` 同口径。

🔴 **表头别名表不翻译**（`_HEADER_ALIASES`）。判据是「拿比较还是显示」：
   它拿来**比较**（判断这一列是哪一列），一旦 `tr()` 过，界面语言一变
   导入就认不出列了。与 `excel_handler.COLUMN_ALIASES`、`_IOT_VENDOR_KEYWORDS`
   同一条规矩。

🔴 **空单元格默认跳过，不等于清空**。这是本模块最要紧的一条语义：
   填表的人留空通常表示「这一格我不知道 / 不填」，若把空当成清空，
   一次导入就会把库里已有的台账**静默抹掉一大片**，而且无法撤销。
   真要批量清空某几列，必须显式传 `clear_empty=True`（CLI 是 `--clear-empty`）。

🔴 **只写生命周期列**。写入一律经 `history_store.update_asset_lifecycle()`，
   它只碰 v7 那六列；扫描写入的列（hostname / vendor / …）一个都不动。
"""
import csv
import datetime
import io
import os
import re

import asset_lifecycle

# ── 行状态（🔴 语言无关的规范值，绝不 tr()）───────────────────────────
#
# `ok` 与 `nochange` 必须分开：把「值本来就一样」算进「已写入」会让用户以为
# 每次导入都改了一堆东西，久了就没人看导入结果了。
ST_OK = "ok"                  # 有实际变化，**将**写入（dry-run 下就停在这里）
ST_WRITTEN = "written"        # 已写入（apply_plan 之后才会出现）
ST_NOCHANGE = "nochange"      # 没有可写字段，或值与库里已一致
ST_NOTFOUND = "notfound"      # 给了身份线索，但库里没有这台设备
ST_INVALID = "invalid"        # 身份缺失 / 字段值非法 / 写入被库层拒绝
ST_DUPLICATE = "duplicate"    # 同一台设备在文件里出现多次（保留第一行）
ST_EMPTY = "empty"            # 整行全空（表格末尾常见，不算错误）

#: 全部状态（顺序即报告顺序：先成功，再跳过，最后失败）。
ALL_STATUSES = (ST_WRITTEN, ST_OK, ST_NOCHANGE, ST_DUPLICATE,
                ST_EMPTY, ST_NOTFOUND, ST_INVALID)

# ── 失败原因（同上，语言无关）────────────────────────────────────────
REASON_BAD_DATE = "bad_date"
REASON_BAD_JSON = "bad_json"
REASON_NO_IDENTITY = "no_identity"
REASON_UNKNOWN_DEVICE = "unknown_device"
REASON_IP_NOT_ALLOWED = "ip_not_allowed"
REASON_DUPLICATE_ROW = "duplicate_row"
REASON_NO_FIELDS = "no_fields"
REASON_ALREADY_EQUAL = "already_equal"
REASON_WRITE_FAILED = "write_failed"

# ── 计划级问题（整张表都用不了，要在动手之前就拦下）──────────────────
PROBLEM_NONE = ""
PROBLEM_NO_COLUMNS = "no_columns"                # 一个可写列都没认出来
PROBLEM_NO_IDENTITY_COLUMN = "no_identity_column"  # 既没身份键列也没 MAC 列
PROBLEM_EMPTY = "empty"                          # 文件里连表头都没有

# ── 表头别名 ─────────────────────────────────────────────────────────
#
# 🔴 这张表**不翻译**（见模块顶部说明）。
# 只收**无歧义**的别名：`位置` / `使用者` 故意**不收** ——
# 资产表里 `location` / `user` 是扫描写出来的列，与人工台账的
# `location_tree` / `owner` 是两回事。把「位置」当「位置树」写进去，
# 等于用扫描到的值覆盖用户手填的台账，正是本模块最该避免的事。
_HEADER_ALIASES = {
    "identity_key": ("identity_key", "身份键"),
    "mac": ("mac", "mac地址", "物理地址", "硬件地址", "mac address"),
    "ip": ("ip", "ip地址", "ip address"),
    "purchase_date": ("purchase_date", "采购日期", "购买日期"),
    "warranty_end": ("warranty_end", "保修到期", "保修截止", "质保到期"),
    "owner": ("owner", "负责人", "责任人"),
    "serial": ("serial", "序列号", "sn"),
    "location_tree": ("location_tree", "位置树"),
    "custom_json": ("custom_json", "自定义字段"),
}

#: 非写入列（只用来认设备 / 只用来报错）。
_IDENTITY_COLUMNS = ("identity_key", "mac", "ip")

_MAC_CLEAN_RE = re.compile(r"[^0-9a-f]")
_HEADER_JUNK_RE = re.compile(r"[\s\*＊]+")


def _norm_header(text) -> str:
    """表头归一：去 BOM、去空白与星号（Excel 的"必填"标记）、转小写。"""
    s = str(text if text is not None else "").lstrip("\ufeff")
    s = _HEADER_JUNK_RE.sub("", s)
    return s.lower()


def _build_alias_map():
    """归一后的表头 → 规范列名。模块加载时建一次，导入时零成本。

    ⚠️ 写成函数而不是模块级 `for` 循环 + `del`：后者会让类型检查报
       "possibly unbound"（循环变量在别名表为空时确实没绑定），
       而为了消警告去加 `# type: ignore` 只会把真问题一起盖掉。
    """
    out = {}
    for col, aliases in _HEADER_ALIASES.items():
        for alias in aliases:
            out[_norm_header(alias)] = col
    return out


_ALIAS_TO_COLUMN = _build_alias_map()


def _mac_key(value) -> str:
    """MAC → 纯小写 hex（去分隔符）；不像 MAC 时返回空串。

    ⚠️ 与 `topology_overrides.normalize_mac()` / `excel_handler.normalize_mac()`
       **刻意重复实现**（那两处也是同样的理由）：本模块要能被 CLI 独立加载，
       不该为了一个 6 行函数去拉起别人的依赖链。三处口径一致性由
       `tests/test_asset_import.py::MacKeyConsistencyTest` 钉住。
    """
    if not value:
        return ""
    raw = _MAC_CLEAN_RE.sub("", str(value).strip().lower())
    return raw if len(raw) == 12 else ""


def _cell_text(value) -> str:
    """单元格值 → 字符串。

    Excel 会把整数读成 `1.0`、日期读成 `datetime` —— 不处理的话
    序列号会变成 `12345.0`、采购日期会变成 `2026-05-01 00:00:00`
    （后者连 `parse_date()` 都认不出来，整行被判非法）。
    """
    if value is None:
        return ""
    # 🔴 bool 是 int 的子类，必须排在 int 之前判断。
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value == int(value) else str(value)
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def detect_columns(headers):
    """表头行 → `(mapping, ignored)`。

    * ``mapping``：``{规范列名: 列索引}``，同一列出现多次时**保留第一处**
      （后面的多半是重复粘贴的列；让后者覆盖前者会悄悄换掉取值来源）；
    * ``ignored``：没认出来的表头原文（去空），**报给用户** ——
      不报的话，用户以为那一列也导进去了，其实是空的。
    """
    mapping, ignored = {}, []
    for idx, raw in enumerate(headers or ()):
        name = _norm_header(raw)
        if not name:
            continue
        col = _ALIAS_TO_COLUMN.get(name)
        if col is None:
            ignored.append(str(raw if raw is not None else "").strip())
            continue
        mapping.setdefault(col, idx)
    return mapping, ignored


def read_table(path):
    """读 CSV / Excel 首个工作表 → ``(headers, body)``（body **不含**表头行）。

    * ``.csv`` / ``.txt`` / ``.tsv``：标准库 `csv`，先按 UTF-8(BOM) 解，
      失败回退 GBK（中文 Windows 上 Excel「另存为 CSV」默认就是 GBK）；
      分隔符用 `csv.Sniffer` 猜（`,` / `;` / 制表符），猜不出按扩展名兜底。
    * ``.xlsx`` / ``.xls``：复用 `excel_handler.read_rows()`（openpyxl/xlrd
      在那里按需引入，本模块顶层不为它付代价）。

    读不了时抛 `ValueError`（调用方转成人话），不返回半截数据 ——
    "读到 3 行然后失败"比"一行都没读"更容易让人误判。
    """
    ext = os.path.splitext(path or "")[1].lower()
    if ext in (".csv", ".txt", ".tsv"):
        text = _read_text(path)
        return _split_csv(text, ext)
    import excel_handler
    raw = excel_handler.read_rows(path)
    rows = [tuple(r) for r in (raw or ())]
    if not rows:
        return (), []
    return tuple(rows[0]), rows[1:]


def _read_text(path) -> str:
    """按 UTF-8(BOM) → GBK 的顺序读文本文件。"""
    last = None
    for enc in ("utf-8-sig", "gbk"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                return f.read()
        except UnicodeDecodeError as exc:
            last = exc
        except OSError as exc:
            raise ValueError(str(exc)) from exc
    raise ValueError("无法解码文件（试过 UTF-8 与 GBK）：%s" % last)


def _split_csv(text, ext):
    """CSV 文本 → ``(headers, body)``。"""
    sample = text[:4096]
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        # 样本太短 / 只有一列 / 引号不配对 —— Sniffer 猜不出来。
        # 这不是错误，按扩展名兜底即可（`except ...: pass` 会被静默吞异常的
        # 棘轮测试拦下，而且这里本来就该给出一个确定的分隔符）。
        delim = "\t" if ext == ".tsv" else ","
    rows = [tuple(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]
    if not rows:
        return (), []
    return rows[0], rows[1:]


def _index_assets(assets):
    """资产列表 → ``(by_key, by_mac, by_ip)`` 三个索引。"""
    by_key, by_mac, by_ip = {}, {}, {}
    for row in assets or ():
        if not isinstance(row, dict):
            continue
        key = (row.get("identity_key") or "").strip()
        if not key:
            continue
        by_key[key] = row
        mac = _mac_key(row.get("mac"))
        if mac and mac not in by_mac:
            by_mac[mac] = key
        ip = (row.get("ip") or "").strip()
        if ip and ip not in by_ip:
            by_ip[ip] = key
    return by_key, by_mac, by_ip


def _item(line, status, reason="", detail="", key="", fields=None,
          field=""):
    """构造一条计划项。键名固定英文 —— CLI 的 `-f json` 直接吐它。"""
    return {
        "line": line,
        "status": status,
        "reason": reason,
        "detail": detail,
        "field": field,
        "identity_key": key,
        "fields": dict(fields or {}),
    }


def build_plan(rows, assets, clear_empty=False, allow_ip=False):
    """把表格行编译成导入计划（**不写库、不改入参**）。

    Args:
        rows:        二维行数据，**第一行是表头**（`assets -f csv` 的产出即符合）。
        assets:      `history_store.list_assets()` 的结果，用来把 MAC / 身份键
                     认回真实身份键。传空列表 ⇒ 所有行都会 `notfound`。
        clear_empty: 空单元格是否视为「清空该字段」。**默认 False = 跳过**。
        allow_ip:    允许只凭 IP 匹配设备。**默认 False** —— IP 会被 DHCP
                     回收后分给另一台机器，按 IP 认设备可能把 A 的负责人
                     写到 B 头上。只有用户明确知道自己在干什么才开。

    Returns:
        dict：``{"problem", "columns", "ignored", "items", "total"}``。
        ``problem`` 非空表示**整张表都用不了**（没有可写列 / 没有认设备的列），
        调用方应当直接报错、别去写。
    """
    rows = [tuple(r or ()) for r in (rows or ())]
    if not rows:
        return {"problem": PROBLEM_EMPTY, "columns": {}, "ignored": [],
                "items": [], "total": 0}

    headers, body = rows[0], rows[1:]
    mapping, ignored = detect_columns(headers)

    plan = {"problem": PROBLEM_NONE, "columns": dict(mapping),
            "ignored": ignored, "items": [], "total": len(body)}

    if not any(c in mapping for c in asset_lifecycle.LIFECYCLE_COLUMNS):
        plan["problem"] = PROBLEM_NO_COLUMNS
        return plan
    # `ip` 也在 `_IDENTITY_COLUMNS` 里：只有 IP 列的表**能**编译出计划，
    # 只是每行会报 `ip_not_allowed` —— 那条信息比「整张表认不出设备」具体得多
    # （它直接告诉用户该加一列 MAC，或显式加 --allow-ip）。
    if not any(c in mapping for c in _IDENTITY_COLUMNS):
        plan["problem"] = PROBLEM_NO_IDENTITY_COLUMN
        return plan

    by_key, by_mac, by_ip = _index_assets(assets)
    seen = {}
    for offset, raw in enumerate(body):
        line = offset + 2  # 表头是第 1 行
        item = _plan_row(raw, line, mapping, by_key, by_mac, by_ip,
                         clear_empty=clear_empty, allow_ip=allow_ip)
        # 重复行：同一台设备在一张表里出现两次，保留第一行、后面的报出来。
        # 静默合并会让用户以为两行的值都生效了（实际只有一行能落库）。
        if item["status"] in (ST_OK, ST_NOCHANGE) and item["identity_key"]:
            prev = seen.get(item["identity_key"])
            if prev is not None:
                item["status"] = ST_DUPLICATE
                item["reason"] = REASON_DUPLICATE_ROW
                item["detail"] = str(prev)
            else:
                seen[item["identity_key"]] = line
        plan["items"].append(item)
    return plan


def _plan_row(raw, line, mapping, by_key, by_mac, by_ip,
              clear_empty=False, allow_ip=False):
    """编译单行。身份解析 → 字段归一与校验 → 与库里的值比对。"""
    cells = list(raw or ())

    def cell(col):
        idx = mapping.get(col)
        if idx is None or idx >= len(cells):
            return ""
        return _cell_text(cells[idx])

    if not any(_cell_text(c) for c in cells):
        return _item(line, ST_EMPTY, REASON_NO_FIELDS)

    raw_key = cell("identity_key").strip()
    mac_text = cell("mac").strip()
    ip_text = cell("ip").strip()

    if raw_key:
        key = raw_key if raw_key in by_key else ""
        if not key:
            return _item(line, ST_NOTFOUND, REASON_UNKNOWN_DEVICE, raw_key)
    elif _mac_key(mac_text):
        key = by_mac.get(_mac_key(mac_text), "")
        if not key:
            return _item(line, ST_NOTFOUND, REASON_UNKNOWN_DEVICE, mac_text)
    elif ip_text and allow_ip:
        key = by_ip.get(ip_text, "")
        if not key:
            return _item(line, ST_NOTFOUND, REASON_UNKNOWN_DEVICE, ip_text)
    elif ip_text:
        # 有 IP 却没开 allow_ip：这不是"找不到设备"，而是"我们拒绝按 IP 认"。
        # 两种原因必须分开报，否则用户会去查库，而真正要做的是加一列 MAC。
        return _item(line, ST_INVALID, REASON_IP_NOT_ALLOWED, ip_text)
    else:
        return _item(line, ST_INVALID, REASON_NO_IDENTITY)

    existing = by_key.get(key) or {}
    fields = {}
    for col in asset_lifecycle.LIFECYCLE_COLUMNS:
        if col not in mapping:
            continue
        text = cell(col)
        if not text:
            if clear_empty:
                fields[col] = ""
            continue
        if col in ("purchase_date", "warranty_end"):
            parsed = asset_lifecycle.parse_date(text)
            if parsed is None:
                return _item(line, ST_INVALID, REASON_BAD_DATE, text,
                             key=key, field=col)
            fields[col] = parsed.isoformat()
        elif col == "custom_json":
            try:
                fields[col] = asset_lifecycle.normalize_custom_json(text)
            except ValueError:
                return _item(line, ST_INVALID, REASON_BAD_JSON, text,
                             key=key, field=col)
        else:
            fields[col] = text

    if not fields:
        return _item(line, ST_NOCHANGE, REASON_NO_FIELDS, key=key)

    # 与库里已有值逐字段比对：全都一样就没必要写，也不算"已导入"。
    if all((existing.get(col) or "") == val for col, val in fields.items()):
        return _item(line, ST_NOCHANGE, REASON_ALREADY_EQUAL, key=key)

    return _item(line, ST_OK, key=key, fields=fields)


def summarize(plan):
    """计划 → ``{状态: 台数}``（只含有出现的状态）。"""
    counts = {}
    for item in (plan or {}).get("items") or ():
        st = item.get("status") or ST_INVALID
        counts[st] = counts.get(st, 0) + 1
    return counts


def pending_count(plan) -> int:
    """还有多少条**待写入**（`ok` 状态）。"""
    return sum(1 for it in (plan or {}).get("items") or ()
               if it.get("status") == ST_OK)


def apply_plan(plan, db_path=None) -> dict:
    """执行计划里所有 `ok` 项。返回 ``{"written", "failed"}``。

    🔴 写入**只走** `history_store.update_asset_lifecycle()` —— 人工台账
       唯一的写入口（它只碰生命周期列，扫描永远覆盖不了）。

    * 成功项状态改为 `written`（所以对同一份计划重复 apply 不会重复写）；
    * `update_asset_lifecycle()` 抛 `ValueError`（库层拒绝）→ 该项改判
      `invalid` / `write_failed` 并计入 `failed`，**其余行继续**；
    * 其它异常（库坏了、磁盘满了）**不吞**，直接抛给调用方 ——
      继续往下写只会产出一份自相矛盾的报告。
    """
    import history_store

    written, failed = 0, []
    for item in (plan or {}).get("items") or ():
        if item.get("status") != ST_OK:
            continue
        try:
            changed = history_store.update_asset_lifecycle(
                item["identity_key"], db_path=db_path, **item["fields"])
        except ValueError as exc:
            item["status"] = ST_INVALID
            item["reason"] = REASON_WRITE_FAILED
            item["detail"] = str(exc)
            failed.append(item)
            continue
        if not changed:
            # 计划编译之后这台设备被清理了（并发场景）。按"找不到"报，
            # 不能算成功 —— 否则用户以为写进去了。
            item["status"] = ST_NOTFOUND
            item["reason"] = REASON_UNKNOWN_DEVICE
            item["detail"] = item["identity_key"]
            failed.append(item)
            continue
        item["status"] = ST_WRITTEN
        written += 1
    return {"written": written, "failed": failed}


# ══════════════════════════════════════════════════════════════════════
#  模板
# ══════════════════════════════════════════════════════════════════════
#
# 为什么需要专门的模板（而不是"拿 `assets -f csv` 的产物当模板"）
# --------------------------------------------------------------
# `assets -f csv` 是 **22 列**：16 个扫描写入列 + 6 个台账列。导入时那 16 列
# 一律被忽略 —— 摆在表里只会让人以为"这些也要填"，填了还没用。
# 模板只保留**导入真正会读**的列：3 个认设备的 + 6 个要填的。
#
# 🔴 模板里**不放示例行**。示例行一旦被漏删就会跟着导入，产生一条
#    `notfound`（示例 MAC 在库里当然不存在）；更糟的是有人照着示例的
#    假数据改了真设备。示例只写在 xlsx 的「说明」页里，数据页永远是干净的。

#: 模板列：认设备的三列 + 六个台账列。顺序即表格里的列顺序。
TEMPLATE_COLUMNS = (("identity_key", "mac", "ip")
                    + tuple(asset_lifecycle.LIFECYCLE_COLUMNS))


def build_template_rows(assets=()):
    """模板的二维数据（**第一行是表头**）。

    给了 `assets` 就把当前资产一台一行铺出来（身份键 / MAC / IP 填好、
    六个台账列留空）—— 这才是实际好用的形态：已有部署直接拿到一张
    「待填清单」，不用自己去查身份键。`assets` 为空时只返回表头。
    """
    rows = [[str(col) for col in TEMPLATE_COLUMNS]]
    blanks = ["" for _ in asset_lifecycle.LIFECYCLE_COLUMNS]
    for asset in assets or ():
        if not isinstance(asset, dict):
            continue
        rows.append([
            (asset.get("identity_key") or "").strip(),
            (asset.get("mac") or "").strip(),
            (asset.get("ip") or "").strip(),
        ] + list(blanks))
    return rows


def write_template(path, assets=(), notes=(), sheet_title="assets",
                   notes_title="readme") -> int:
    """写模板文件，返回**数据行数**（不含表头）。

    * ``.csv``：UTF-8 **带 BOM**（Excel 双击不乱码），只写数据页；
      `notes` 会被忽略 —— CSV 装不下第二个 sheet，硬塞进去会变成一行脏数据。
    * ``.xlsx``：数据页 + 说明页（`notes` 逐行写进去，自动换行）。

    ⚠️ 本模块**零 i18n**，所以 `notes` / `sheet_title` / `notes_title`
       都由调用方给**已翻译好**的文本；不传就用英文默认值。
    """
    ext = os.path.splitext(path or "")[1].lower()
    rows = build_template_rows(assets)
    if ext in (".csv", ".txt"):
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f, lineterminator="\n").writerows(rows)
        return len(rows) - 1
    if ext == ".xlsx":
        _write_xlsx_template(path, rows, notes, sheet_title, notes_title)
        return len(rows) - 1
    raise ValueError("模板只支持 .csv 或 .xlsx（收到 %r）"
                     % (ext or "无扩展名"))


def _write_xlsx_template(path, rows, notes, sheet_title, notes_title):
    """写 xlsx：数据页（带表头样式）+ 说明页。"""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    # 🔴 用 `worksheets[0]` 而不是 `wb.active`：后者类型上是 `Worksheet | None`
    #    （类型检查会要求每处都判空），而新建的工作簿**必然**恰好有一张表。
    ws = wb.worksheets[0]
    ws.title = sheet_title

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB",
                              fill_type="solid")
    centered = Alignment(horizontal="center", vertical="center")
    border = Border(left=Side(style="thin", color="D1D5DB"),
                    right=Side(style="thin", color="D1D5DB"),
                    top=Side(style="thin", color="D1D5DB"),
                    bottom=Side(style="thin", color="D1D5DB"))

    for row_idx, values in enumerate(rows, 1):
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = border
            if row_idx == 1:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = centered
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    widths = {"identity_key": 26, "mac": 20, "ip": 16, "purchase_date": 14,
              "warranty_end": 14, "owner": 12, "serial": 18,
              "location_tree": 20, "custom_json": 28}
    for col_idx, name in enumerate(TEMPLATE_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(name, 18)
    # 冻结表头：往下翻时列名一直看得见（填几十行时这个很要紧）。
    ws.freeze_panes = "A2"

    if notes:
        ns = wb.create_sheet(notes_title)
        ns.column_dimensions["A"].width = 100
        for row_idx, line in enumerate(notes, 1):
            cell = ns.cell(row=row_idx, column=1, value=line)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)
    wb.close()
