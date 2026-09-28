# -*- coding: utf-8 -*-
"""
MAC 地址库 (持久化)

基于 SQLite，落盘保存，支持：
- Excel 批量合并导入 (自动识别列、按 MAC/IP 去重、空字段回填)
- 手动 增加 / 修改 / 删除 单条
- 按 IP 或 MAC 精确匹配 (扫描时自动标注设备类型/位置/使用者)
- 导出为 Excel
- 重启后自动加载 (数据库文件默认 mac_library.db)
"""
import contextlib
import os
from datetime import datetime
from typing import List, Optional, Tuple
from app_paths import data_dir
from app_logging import get_logger
from i18n import tr

_log = get_logger(__name__)

import db_util
from excel_handler import (
    normalize_mac,
    detect_columns,
    _read_rows,
    _cell_to_str,
    COLUMN_ALIASES,
)

DEFAULT_DB = os.path.join(data_dir(), "mac_library.db")

# 基礎表结构 + 非唯一索引；部分唯一索引在 _init_db 里于去重后创建
_SCHEMA_BASE = """
CREATE TABLE IF NOT EXISTS devices (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    ip          TEXT DEFAULT '',
    mac         TEXT DEFAULT '',
    device_type TEXT DEFAULT '',
    location    TEXT DEFAULT '',
    user        TEXT DEFAULT '',
    remark      TEXT DEFAULT '',
    vendor      TEXT DEFAULT '',
    source_file TEXT DEFAULT '',
    created_at  TEXT,
    updated_at  TEXT
);
CREATE INDEX IF NOT EXISTS idx_dev_mac ON devices(mac);
CREATE INDEX IF NOT EXISTS idx_dev_ip  ON devices(ip);
"""

# 批量导入时的"只填空字段、不覆盖已有值"回填语句。
#
# 刻意写成**固定形状**（每条 UPDATE 的 SET 子句完全相同，只有参数不同），
# 这样上千条回填可以走同一次 executemany，而不是每条一个独立 SQL + commit。
# 语义等价于旧实现的 `fill[f] = new_v if new_v and not existing.get(f)`：
#   IFNULL(col,'')='' → 该字段为空才写入新值，否则保持原值。
# 传入空串时对空字段是空操作，因此无需在 Python 侧先过滤。
# 参数顺序：ip, device_type, location, user, remark, vendor, mac, updated_at, id
_FILL_UPDATE_SQL = (
    "UPDATE devices SET "
    "  ip=CASE WHEN IFNULL(ip,'')='' THEN ? ELSE ip END, "
    "  device_type=CASE WHEN IFNULL(device_type,'')='' THEN ? ELSE device_type END, "
    "  location=CASE WHEN IFNULL(location,'')='' THEN ? ELSE location END, "
    "  user=CASE WHEN IFNULL(user,'')='' THEN ? ELSE user END, "
    "  remark=CASE WHEN IFNULL(remark,'')='' THEN ? ELSE remark END, "
    "  vendor=CASE WHEN IFNULL(vendor,'')='' THEN ? ELSE vendor END, "
    "  mac=CASE WHEN IFNULL(mac,'')='' THEN ? ELSE mac END, "
    "  updated_at=? "
    "WHERE id=?"
)


class MacLibrary:
    """落盘的 MAC 地址库"""

    def __init__(self, db_path: str = DEFAULT_DB):
        self.db_path = db_path
        with self._conn():  # 触发建表
            pass

    # ---------- 底层 ----------
    @contextlib.contextmanager
    def _conn(self):
        """复用进程级单例连接（按 db_path 缓存，线程安全）。"""
        with db_util.connection(self.db_path, init_fn=self._init_db) as c:
            yield c

    def _init_db(self, conn):
        # 1) 建表与基础索引
        conn.executescript(_SCHEMA_BASE)
        # 2) 对已有非空 MAC 去重（保留最大 id 即最后写入一条），否则部分唯一索引会创建失败
        try:
            conn.execute(
                "DELETE FROM devices WHERE mac <> '' "
                "AND id NOT IN (SELECT MAX(id) FROM devices WHERE mac <> '' GROUP BY mac)"
            )
        except Exception as e:  # noqa: BLE001
            # 去重失败会导致下一步的唯一索引建不起来，进而失去"MAC 唯一"保证，
            # 因此必须留痕（不中断启动，但要能在日志里看到）
            _log.warning("地址库 MAC 去重失败，唯一索引可能无法建立: %s", e, exc_info=True)
        # 3) 部分唯一索引：仅对非空 MAC 约束唯一（空 MAC 允许多条，用于无 MAC 的占位设备）
        try:
            conn.execute(
                "CREATE UNIQUE INDEX IF NOT EXISTS idx_dev_mac_uniq "
                "ON devices(mac) WHERE mac <> ''"
            )
        except Exception as e:  # noqa: BLE001
            _log.warning("创建 MAC 唯一索引失败，地址库可能出现重复 MAC: %s", e, exc_info=True)
        conn.commit()

    @staticmethod
    def _norm_mac(mac: str) -> str:
        return normalize_mac(mac)

    # ---------- 增删改 ----------
    def add(self, ip="", mac="", device_type="", location="", user="",
            remark="", vendor="", source_file="") -> int:
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        mac = self._norm_mac(mac)
        with self._conn() as conn:
            # 非空 MAC 已存在 → 整行更新（部分唯一索引保护，避免 IntegrityError）
            if mac:
                cur = conn.execute("SELECT id FROM devices WHERE mac=?", (mac,)).fetchone()
                if cur:
                    conn.execute(
                        "UPDATE devices SET ip=?, device_type=?, location=?, user=?, "
                        "remark=?, vendor=?, source_file=?, updated_at=? WHERE id=?",
                        (ip, device_type, location, user, remark, vendor, source_file, now, cur[0]),
                    )
                    conn.commit()
                    return cur[0]
            cur = conn.execute(
                "INSERT INTO devices "
                "(ip, mac, device_type, location, user, remark, vendor, source_file, created_at, updated_at) "
                "VALUES (?,?,?,?,?,?,?,?,?,?)",
                (ip, mac, device_type, location, user, remark, vendor, source_file, now, now),
            )
            conn.commit()
            return cur.lastrowid

    _EDITABLE = {"ip", "mac", "device_type", "location", "user", "remark", "vendor", "source_file"}

    def update(self, device_id: int, **fields):
        sets, vals = [], []
        for k, v in fields.items():
            if k in self._EDITABLE:
                val = self._norm_mac(v) if k == "mac" else v
                sets.append(f"{k}=?")
                vals.append(val)
        if not sets:
            return
        sets.append("updated_at=?")
        vals.append(datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        vals.append(device_id)
        with self._conn() as conn:
            conn.execute(f"UPDATE devices SET {', '.join(sets)} WHERE id=?", vals)
            conn.commit()

    def delete(self, device_id: int):
        with self._conn() as conn:
            conn.execute("DELETE FROM devices WHERE id=?", (device_id,))
            conn.commit()

    # ---------- 查询 ----------
    def get_all(self) -> List[dict]:
        with self._conn() as conn:
            rows = conn.execute("SELECT * FROM devices ORDER BY id").fetchall()
        return [dict(r) for r in rows]

    def get_by_id(self, device_id: int) -> Optional[dict]:
        with self._conn() as conn:
            r = conn.execute("SELECT * FROM devices WHERE id=?", (device_id,)).fetchone()
        return dict(r) if r else None

    def search(self, query: str) -> List[dict]:
        q = f"%{query.strip()}%"
        with self._conn() as conn:
            rows = conn.execute(
                "SELECT * FROM devices WHERE ip LIKE ? OR mac LIKE ? OR device_type LIKE ? "
                "OR location LIKE ? OR user LIKE ? OR remark LIKE ? OR vendor LIKE ? "
                "ORDER BY id",
                (q, q, q, q, q, q, q),
            ).fetchall()
        return [dict(r) for r in rows]

    def count(self) -> int:
        with self._conn() as conn:
            c = conn.execute("SELECT COUNT(*) FROM devices").fetchone()[0]
        return c

    def source_files(self) -> List[str]:
        with self._conn() as conn:
            rows = conn.execute(
                "SELECT DISTINCT source_file FROM devices WHERE source_file != ''"
            ).fetchall()
        return [r["source_file"] for r in rows]

    def clear(self):
        with self._conn() as conn:
            conn.execute("DELETE FROM devices")
            conn.commit()

    # ---------- 匹配 ----------
    def match(self, ip: str, mac: str = "") -> Optional[dict]:
        """优先 IP 精确匹配，其次 MAC 匹配。返回设备 dict 或 None。"""
        with self._conn() as conn:
            if ip:
                r = conn.execute("SELECT * FROM devices WHERE ip=?", (ip.strip(),)).fetchone()
                if r:
                    return dict(r)
            mac_norm = self._norm_mac(mac)
            if mac_norm:
                r = conn.execute("SELECT * FROM devices WHERE mac=?", (mac_norm,)).fetchone()
                if r:
                    return dict(r)
        return None

    def is_loaded(self) -> bool:
        return self.count() > 0

    # ---------- Excel 合并导入 ----------
    def merge_from_excel(self, file_path: str) -> Tuple[bool, str]:
        """
        解析 Excel 并合并进地址库。
        去重规则：按 MAC(优先) 或 IP 判断已存在；若已存在则只回填空字段，不覆盖已有值。
        """
        try:
            rows = _read_rows(file_path)
        except Exception as e:
            return False, f"无法打开文件: {e}"
        if not rows:
            return False, tr("文件为空")

        # 表头行定位 (兼容首行是大标题的情况)
        header_row_idx = 0
        for idx in range(min(15, len(rows))):
            hdr = [str(h).strip().lower() if h is not None else "" for h in rows[idx]]
            if any(any(h in aliases for h in hdr) for aliases in COLUMN_ALIASES.values()):
                header_row_idx = idx
                break

        raw_headers = [str(h).strip() if h is not None else "" for h in rows[header_row_idx]]
        mapping = detect_columns(raw_headers)
        if "ip" not in mapping and "mac" not in mapping:
            return False, tr("未找到 IP/MAC 列，无法自动匹配设备（请在清单中补充 MAC 或 IP 列后重试）")

        ip_cols = mapping.get("ip", [])
        mac_cols = mapping.get("mac", [])
        added = updated = skipped = 0
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        src = os.path.basename(file_path)

        # 导入列（与 add() 的 INSERT 列顺序一致）
        insert_cols = ("ip", "mac", "device_type", "location", "user",
                       "remark", "vendor")

        with self._conn() as conn:
            # 一次性把已有数据读进内存建索引，替代逐行 self.match()。
            #
            # 旧写法每行要 1~2 次 SELECT（match）+ 1 次写 + 1 次 commit，
            # 1000 行的清单就是 3000+ 次查询和 1000 个独立事务。
            # 现在：1 次全量读 + 内存匹配 + 2 次批量写 + 1 次 commit。
            #
            # ORDER BY id ASC 是为了让 setdefault 的"先到先得"与旧 match()
            # 的 fetchone()（无 ORDER BY，实际取最低 rowid）语义一致。
            by_ip, by_mac = {}, {}
            for r in conn.execute("SELECT * FROM devices ORDER BY id ASC").fetchall():
                d = dict(r)
                if d.get("ip"):
                    by_ip.setdefault(str(d["ip"]).strip(), d)
                if d.get("mac"):
                    by_mac.setdefault(d["mac"], d)

            pending = []   # 本次待插入的新行（dict）
            updates = []   # 待回填的已有行参数

            for row in rows[header_row_idx + 1:]:
                if not row or all(c is None or str(c).strip() == "" for c in row):
                    continue

                entry = {}
                for field, cols in mapping.items():
                    vals = []
                    for c in cols:
                        if c < len(row):
                            v = _cell_to_str(row[c])
                            if v:
                                vals.append(v)
                    if vals:
                        entry[field] = " ".join(vals)

                mac_raw = ""
                for c in mac_cols:
                    if c < len(row):
                        m = _cell_to_str(row[c])
                        if m:
                            mac_raw = m
                            break

                ip_val = ""
                if ip_cols and ip_cols[0] < len(row):
                    ip_val = _cell_to_str(row[ip_cols[0]]).strip()

                if not ip_val and not mac_raw:
                    continue

                mac_norm = self._norm_mac(mac_raw)

                # 内存内匹配：优先 IP 精确，其次 MAC
                existing = by_ip.get(ip_val) if ip_val else None
                if existing is None and mac_norm:
                    existing = by_mac.get(mac_norm)

                if existing is not None:
                    fill = {}
                    for f in ("ip", "device_type", "location", "user", "remark", "vendor"):
                        new_v = entry.get(f, "")
                        if new_v and not existing.get(f):
                            fill[f] = new_v
                    if mac_norm and not existing.get("mac"):
                        fill["mac"] = mac_norm

                    if not fill:
                        skipped += 1
                    else:
                        if "id" in existing:
                            # 已在库中：记一条批量回填
                            updates.append(tuple(
                                fill.get(f, "") for f in
                                ("ip", "device_type", "location", "user",
                                 "remark", "vendor", "mac")
                            ) + (now, existing["id"]))
                        # 同步内存索引：同一文件里后面再出现同一设备时，
                        # 应看到刚补上的值（与旧实现逐行读库的效果一致）。
                        existing.update(fill)
                        updated += 1
                else:
                    rec = {
                        "ip": ip_val,
                        "mac": mac_norm,
                        "device_type": entry.get("device_type", ""),
                        "location": entry.get("location", ""),
                        "user": entry.get("user", ""),
                        "remark": entry.get("remark", ""),
                        "vendor": entry.get("vendor", ""),
                    }
                    pending.append(rec)
                    added += 1
                    # 登记进索引，使同文件内的重复行按"已存在"处理
                    if rec["ip"]:
                        by_ip.setdefault(rec["ip"], rec)
                    if rec["mac"]:
                        by_mac.setdefault(rec["mac"], rec)

            # 批量落库。整体放在一个事务里：要么全部成功，要么一行不落。
            # 旧实现是逐行 commit，中途失败会留下"导了一半"的库 —— 用户重试时
            # 又分不清哪些已导入。
            try:
                if pending:
                    conn.executemany(
                        "INSERT INTO devices "
                        "(ip, mac, device_type, location, user, remark, vendor, "
                        " source_file, created_at, updated_at) "
                        "VALUES (?,?,?,?,?,?,?,?,?,?)",
                        [tuple(rec[c] for c in insert_cols) + (src, now, now)
                         for rec in pending],
                    )
                if updates:
                    conn.executemany(_FILL_UPDATE_SQL, updates)
                if pending or updates:
                    conn.commit()
            except Exception as e:  # noqa: BLE001
                conn.rollback()
                # 最可能的成因：同一 MAC 被清单里两行分别指给了不同 IP，
                # 回填时撞上 idx_dev_mac_uniq（非空 MAC 唯一）。
                _log.warning("合并地址库失败，已回滚本次导入: %s", e, exc_info=True)
                return False, f"合并失败（已回滚，地址库未改动）: {e}"

        if added == 0 and updated == 0 and skipped == 0:
            return False, tr("MAC/IP 数据为空，该文件无法在扫描时自动匹配设备（请检查 MAC 列是否已填写）")

        msg = (f"已从 {os.path.basename(file_path)} 合并：新增 {added} 条，"
               f"更新已有 {updated} 条，无变化跳过 {skipped} 条")
        return True, msg

    # ---------- 导出 ----------
    def export_to_excel(self, file_path: str) -> Tuple[bool, str]:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = tr("MAC地址库")
            headers = ["ID", tr("IP地址"), tr("MAC地址"), tr("设备类型"), tr("位置"), tr("使用者"), tr("备注"), tr("来源文件")]
            hf = Font(name=tr("微软雅黑"), size=11, bold=True, color="FFFFFF")
            hf_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
            for ci, h in enumerate(headers, 1):
                c = ws.cell(row=1, column=ci, value=h)
                c.font = hf
                c.fill = hf_fill
                c.alignment = Alignment(horizontal="center", vertical="center")

            devices = self.get_all()
            for ri, d in enumerate(devices, 2):
                vals = [d["id"], d["ip"], d["mac"], d["device_type"], d["location"],
                        d["user"], d["remark"], d["source_file"]]
                for ci, v in enumerate(vals, 1):
                    ws.cell(row=ri, column=ci, value=v)

            widths = [6, 16, 20, 14, 22, 14, 22, 30]
            for ci, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(ci)].width = w
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
            wb.save(file_path)
            wb.close()
            return True, f"已导出 {len(devices)} 条到 {file_path}"
        except Exception as e:
            return False, f"导出失败: {e}"


class MacMatcher:
    """地址库内存索引：把整库预加载为 {ip: row} / {mac: row}，做 O(1) 匹配。

    为什么需要它（P0-3）：
    原先每扫到一台设备就在 GUI 主线程查一次 SQLite（/16 场景约 6.5 万次查询，
    界面会长时间无响应）。地址库通常只有几千条，一次性读进内存后匹配就是
    纯字典查找，而且可以安全地在扫描工作线程内完成，主线程只负责渲染。

    - 匹配语义与 MacLibrary.match 保持一致：优先 IP 精确匹配，其次 MAC 匹配。
    - 同一 IP / MAC 命中多行时取 id 最小的一行（get_all 按 id 升序，先到先得）。
    - 纯内存对象，无 SQLite 依赖，可在任意线程使用。
    """

    __slots__ = ("_by_ip", "_by_mac", "size")

    def __init__(self, rows=None):
        self._by_ip = {}
        self._by_mac = {}
        for d in rows or ():
            ip = (d.get("ip") or "").strip()
            if ip and ip not in self._by_ip:
                self._by_ip[ip] = d
            mac = normalize_mac(d.get("mac") or "")
            if mac and mac not in self._by_mac:
                self._by_mac[mac] = d
        self.size = len(self._by_ip) or len(self._by_mac)

    @classmethod
    def from_library(cls, library):
        """从 MacLibrary 实例构建（一次 SELECT 读全表）。

        库为空 / 未提供时返回 None，调用方据此跳过匹配。
        """
        if library is None:
            return None
        try:
            if not library.is_loaded():
                return None
            return cls(library.get_all())
        except Exception:
            return None

    def match(self, ip: str, mac: str = ""):
        """返回命中的设备 dict；未命中返回 None。"""
        row = self._by_ip.get((ip or "").strip())
        if row is not None:
            return row
        m = normalize_mac(mac or "")
        if m:
            return self._by_mac.get(m)
        return None

    def apply_to(self, result) -> bool:
        """把命中的台账字段写入 ScanResult；命中返回 True。

        未命中时不做任何改动（不会清空已有标注），与旧行为一致。
        """
        row = self.match(getattr(result, "ip", ""), getattr(result, "mac", ""))
        if not row:
            return False
        result.device_type = row.get("device_type", "") or ""
        result.location = row.get("location", "") or ""
        result.user = row.get("user", "") or ""
        result.remark = row.get("remark", "") or ""
        result.excel_matched = True
        return True
